from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, PP_ALIGN as _PA
from pptx.enum.shapes import MSO_SHAPE
import os

OUTPUT_LOCAL = "从小微专营到网点变革.pptx"
OUTPUT_TARGET = r"E:\download\泰隆\泰隆ppt\从小微专营到网点变革.pptx"

ORANGE = RGBColor(0xF2, 0x6C, 0x21)
ORANGE_DARK = RGBColor(0xD3, 0x54, 0x00)
ORANGE_LIGHT = RGBColor(0xFF, 0xE8, 0xD6)
ORANGE_SOFT = RGBColor(0xFF, 0xF3, 0xEA)
TEXT_DARK = RGBColor(0x2B, 0x2B, 0x2B)
TEXT_GRAY = RGBColor(0x6B, 0x6B, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG_CREAM = RGBColor(0xFF, 0xFB, 0xF7)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height

from pptx.enum.text import MSO_ANCHOR


def add_rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    return s


def add_text(slide, x, y, w, h, text, size=18, bold=False, color=TEXT_DARK,
             align=PP_ALIGN.LEFT, font="微软雅黑", anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return tb


def add_page_number(slide, num, total):
    add_text(slide, Inches(12.3), Inches(7.05), Inches(0.9), Inches(0.35),
             "%d / %d" % (num, total), size=10, color=TEXT_GRAY, align=PP_ALIGN.RIGHT)


def add_top_bar(slide):
    add_rect(slide, 0, 0, SW, Inches(0.08), fill=ORANGE)


def add_bottom_bar(slide):
    add_rect(slide, 0, Inches(7.42), SW, Inches(0.08), fill=ORANGE_LIGHT)


def img_placeholder(slide, x, y, w, h, label="此处插入图片"):
    add_rect(slide, x, y, w, h, fill=ORANGE_SOFT, line=ORANGE_LIGHT)
    add_text(slide, x, y + h / 2 - Inches(0.3), w, Inches(0.6),
             "[图] %s" % label, size=12, color=ORANGE_DARK, align=PP_ALIGN.CENTER)


TOTAL_SLIDES = 15


def slide_cover():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SW, SH, fill=BG_CREAM)
    add_rect(s, 0, 0, Inches(4.8), SH, fill=ORANGE)
    deco = [
        (Inches(4.8), Inches(0), Inches(0.6), Inches(1.5), ORANGE_LIGHT),
        (Inches(4.8), Inches(1.5), Inches(1.1), Inches(0.8), ORANGE_DARK),
        (Inches(4.8), Inches(5.5), Inches(0.9), Inches(1.0), ORANGE_LIGHT),
        (Inches(4.8), Inches(6.5), Inches(0.5), Inches(1.0), ORANGE_DARK),
    ]
    for x, y, w, h, c in deco:
        add_rect(s, x, y, w, h, fill=c)
    for i, r in enumerate([(Inches(0.6), 0.45), (Inches(0.9), 0.3), (Inches(1.2), 0.18)]):
        fill = WHITE if i % 2 == 0 else ORANGE_LIGHT
        add_rect(s, r[0], Inches(6.2), Inches(r[1]), Inches(r[1]), fill=fill,
                 shape=MSO_SHAPE.OVAL)
    add_text(s, Inches(0.8), Inches(1.8), Inches(3.6), Inches(0.5),
             "INTERNSHIP DEFENSE", size=14, bold=True, color=WHITE, font="Arial")
    add_text(s, Inches(0.8), Inches(2.4), Inches(3.8), Inches(2.0),
             "小微专营\n的网点变革", size=44, bold=True, color=WHITE, font="微软雅黑")
    add_text(s, Inches(0.8), Inches(4.7), Inches(3.6), Inches(0.6),
             "泰隆银行实习结业答辩", size=18, color=ORANGE_LIGHT)
    add_text(s, Inches(5.8), Inches(1.6), Inches(7), Inches(0.5),
             "答辩人 · 曾纪元", size=20, bold=True, color=TEXT_DARK)
    add_text(s, Inches(5.8), Inches(2.2), Inches(7), Inches(0.4),
             "指导人：XXX    |    2026.09", size=14, color=TEXT_GRAY)
    add_rect(s, Inches(5.8), Inches(3.0), Inches(0.8), Inches(0.06), fill=ORANGE)
    items = [
        ("01", "网点现状", '从「金融网点」到「社区节点」'),
        ("02", "需求观察", '乡镇网点的优势不是「流量」，而是「关系」'),
        ("03", "变革方向", '打造「小而实用」的社区中心'),
        ("04", "核心逻辑", '非金融服务带来「人」，金融服务留下「关系」'),
    ]
    yy = Inches(3.4)
    for num, title, sub in items:
        add_text(s, Inches(5.8), yy, Inches(0.8), Inches(0.6),
                 num, size=28, bold=True, color=ORANGE, font="Arial")
        add_text(s, Inches(6.7), yy + Inches(0.02), Inches(6), Inches(0.4),
                 title, size=18, bold=True, color=TEXT_DARK)
        add_text(s, Inches(6.7), yy + Inches(0.42), Inches(6), Inches(0.4),
                 sub, size=12, color=TEXT_GRAY)
        yy += Inches(0.95)
    img_placeholder(s, Inches(8.2), Inches(6.0), Inches(4.5), Inches(1.1),
                    "插入：泰隆网点门头 / 实习合影（横版）")
    add_text(s, Inches(0.8), Inches(7.05), Inches(4), Inches(0.3),
             "TAILONG BANK · 2026", size=10, color=ORANGE_LIGHT)


def slide_toc():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SW, SH, fill=BG_CREAM)
    add_top_bar(s)
    add_bottom_bar(s)
    add_text(s, Inches(0.8), Inches(0.5), Inches(3), Inches(0.4),
             "CONTENTS", size=14, bold=True, color=ORANGE, font="Arial")
    add_text(s, Inches(0.8), Inches(0.95), Inches(6), Inches(0.8),
             "目  录", size=36, bold=True, color=TEXT_DARK)
    add_rect(s, Inches(0.8), Inches(1.85), Inches(1.2), Inches(0.06), fill=ORANGE)
    cards = [
        ("01", "网点现状", "从金融网点到社区节点的角色转变", Inches(0.8)),
        ("02", "需求观察", "流量不是核心，关系才是护城河", Inches(4.1)),
        ("03", "变革方向", "打造小而实用的乡镇社区中心", Inches(7.4)),
        ("04", "核心逻辑", "非金融引人，金融留关系", Inches(10.7)),
    ]
    for num, title, sub, cx in cards:
        add_rect(s, cx, Inches(2.5), Inches(2.8), Inches(4.2), fill=WHITE, line=ORANGE_LIGHT)
        add_rect(s, cx, Inches(2.5), Inches(2.8), Inches(0.12), fill=ORANGE)
        add_text(s, cx + Inches(0.3), Inches(2.9), Inches(2.2), Inches(0.8),
                 num, size=44, bold=True, color=ORANGE, font="Arial")
        add_text(s, cx + Inches(0.3), Inches(3.9), Inches(2.2), Inches(0.6),
                 title, size=22, bold=True, color=TEXT_DARK)
        add_rect(s, cx + Inches(0.3), Inches(4.6), Inches(0.6), Inches(0.05), fill=ORANGE_LIGHT)
        add_text(s, cx + Inches(0.3), Inches(4.85), Inches(2.2), Inches(1.3),
                 sub, size=13, color=TEXT_GRAY)
    add_page_number(s, 2, TOTAL_SLIDES)


def slide_part_intro(part_num, part_cn, part_en, quote, page_num):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SW, SH, fill=ORANGE)
    add_rect(s, Inches(9.5), 0, Inches(3.9), SH, fill=ORANGE_DARK)
    for i, sz in enumerate([3.2, 2.2, 1.4, 0.8]):
        fill = ORANGE_LIGHT if i % 2 == 0 else WHITE
        add_rect(s, Inches(11.5), Inches(5.8 - i * 0.15), Inches(sz * 0.2), Inches(sz * 0.2),
                 fill=fill, shape=MSO_SHAPE.OVAL)
    add_text(s, Inches(0.8), Inches(1.0), Inches(8), Inches(0.5),
             "PART 0%d" % part_num, size=20, bold=True, color=ORANGE_LIGHT, font="Arial")
    add_text(s, Inches(0.8), Inches(1.6), Inches(9), Inches(2.2),
             part_cn, size=52, bold=True, color=WHITE)
    add_rect(s, Inches(0.8), Inches(4.1), Inches(1.5), Inches(0.08), fill=WHITE)
    add_text(s, Inches(0.8), Inches(4.4), Inches(9), Inches(0.6),
             part_en, size=18, color=ORANGE_LIGHT, font="微软雅黑")
    add_text(s, Inches(0.8), Inches(5.5), Inches(8.5), Inches(1.2),
             quote, size=16, color=WHITE, font="楷体")
    add_text(s, Inches(9.8), Inches(2.5), Inches(3.2), Inches(2),
             "0%d" % part_num, size=180, bold=True, color=ORANGE, font="Arial")
    add_page_number(s, page_num, TOTAL_SLIDES)


def slide_p1_status():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SW, SH, fill=BG_CREAM)
    add_top_bar(s)
    add_bottom_bar(s)
    add_text(s, Inches(0.8), Inches(0.45), Inches(2), Inches(0.35),
             "PART 01 · 网点现状", size=12, bold=True, color=ORANGE)
    add_text(s, Inches(0.8), Inches(0.85), Inches(10), Inches(0.8),
             '传统网点的角色：一座「金融孤岛」', size=28, bold=True, color=TEXT_DARK)
    add_rect(s, Inches(0.8), Inches(1.72), Inches(1), Inches(0.06), fill=ORANGE)
    old_items = [
        ("🏦", "物理边界清晰", "玻璃隔断 · 叫号排队 · 柜台是唯一接触面"),
        ("⏰", "服务时间受限", "朝九晚五 · 工作日营业 · 与村民作息错配"),
        ("📊", "产品导向为主", "存款贷款指标驱动 · 完成任务式沟通"),
        ("🚪", "进出动机单一", "只有「办事」才进门 · 办完即走"),
    ]
    yy = Inches(2.1)
    for icon, title, sub in old_items:
        add_rect(s, Inches(0.8), yy, Inches(5.8), Inches(1.0), fill=WHITE, line=ORANGE_LIGHT)
        add_text(s, Inches(1.05), yy + Inches(0.18), Inches(0.7), Inches(0.7),
                 icon, size=26, align=PP_ALIGN.CENTER)
        add_text(s, Inches(1.9), yy + Inches(0.12), Inches(4.5), Inches(0.4),
                 title, size=15, bold=True, color=TEXT_DARK)
        add_text(s, Inches(1.9), yy + Inches(0.52), Inches(4.5), Inches(0.4),
                 sub, size=11, color=TEXT_GRAY)
        yy += Inches(1.15)
    add_rect(s, Inches(7.0), Inches(2.1), Inches(5.5), Inches(4.9), fill=ORANGE_SOFT, line=ORANGE_LIGHT)
    add_text(s, Inches(7.3), Inches(2.3), Inches(5), Inches(0.5),
             '➡  我在乡镇网点观察到的日常', size=14, bold=True, color=ORANGE_DARK)
    img_placeholder(s, Inches(7.3), Inches(2.9), Inches(4.9), Inches(2.4),
                    "插入：网点大厅 / 柜台 / 老人排队实景照片")
    add_text(s, Inches(7.3), Inches(5.5), Inches(4.9), Inches(1.3),
             "· 上午 8:30 门口已聚集老人等候开门\n· 取养老金 = 每周固定「社交日程」\n· 客户经理一出门，整条街都跟他打招呼",
             size=12, color=TEXT_DARK)
    add_page_number(s, 4, TOTAL_SLIDES)


def slide_p1_shift():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SW, SH, fill=BG_CREAM)
    add_top_bar(s)
    add_bottom_bar(s)
    add_text(s, Inches(0.8), Inches(0.45), Inches(2), Inches(0.35),
             "PART 01 · 网点现状", size=12, bold=True, color=ORANGE)
    add_text(s, Inches(0.8), Inches(0.85), Inches(12), Inches(0.8),
             '角色跃迁：从「金融网点」 → 「社区节点」', size=28, bold=True, color=TEXT_DARK)
    add_rect(s, Inches(0.8), Inches(1.72), Inches(1), Inches(0.06), fill=ORANGE)
    arrow_x = Inches(6.35)
    add_rect(s, Inches(0.8), Inches(2.1), Inches(5.2), Inches(4.8), fill=WHITE, line=ORANGE_LIGHT)
    add_rect(s, Inches(0.8), Inches(2.1), Inches(5.2), Inches(0.6), fill=ORANGE_LIGHT)
    add_text(s, Inches(0.8), Inches(2.18), Inches(5.2), Inches(0.45),
             "BEFORE · 金融网点", size=15, bold=True, color=ORANGE_DARK, align=PP_ALIGN.CENTER)
    before = ["办理存取款", "发放贷款", "推销理财产品", "解答账户问题", "—— 纯金融属性"]
    for i, t in enumerate(before):
        col = TEXT_GRAY if i == 4 else TEXT_DARK
        add_text(s, Inches(1.1), Inches(3.0 + i * 0.75), Inches(4.6), Inches(0.5),
                 "▸  %s" % t, size=14, color=col)
    add_text(s, arrow_x, Inches(4.1), Inches(0.6), Inches(0.8),
             "➜", size=40, color=ORANGE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Inches(7.2), Inches(2.1), Inches(5.3), Inches(4.8), fill=ORANGE)
    add_rect(s, Inches(7.2), Inches(2.1), Inches(5.3), Inches(0.6), fill=ORANGE_DARK)
    add_text(s, Inches(7.2), Inches(2.18), Inches(5.3), Inches(0.45),
             "AFTER · 社区节点", size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    after = ["代缴水电费 / 话费", "快递代收点", "避暑歇脚 + 免费茶水", "老人手机课堂", '村口的「信息中转站」']
    for i, t in enumerate(after):
        add_text(s, Inches(7.5), Inches(3.0 + i * 0.75), Inches(4.7), Inches(0.5),
                 "✦  %s" % t, size=14, color=WHITE)
    add_text(s, Inches(0.8), Inches(7.0), Inches(12), Inches(0.35),
             '结论：当网点承担的非金融功能越多，它在社区中的「位置感」就越强。',
             size=12, bold=True, color=ORANGE_DARK, align=PP_ALIGN.CENTER)
    add_page_number(s, 5, TOTAL_SLIDES)


def slide_p2_observation():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SW, SH, fill=BG_CREAM)
    add_top_bar(s)
    add_bottom_bar(s)
    add_text(s, Inches(0.8), Inches(0.45), Inches(2), Inches(0.35),
             "PART 02 · 需求观察", size=12, bold=True, color=ORANGE)
    add_text(s, Inches(0.8), Inches(0.85), Inches(12), Inches(0.8),
             "三组真实对话：我在走访中听到的声音", size=28, bold=True, color=TEXT_DARK)
    add_rect(s, Inches(0.8), Inches(1.72), Inches(1), Inches(0.06), fill=ORANGE)
    cases = [
        ("案例 01", "村口杂货店王阿姨",
         '「钱放在你们这里放心，因为小周上次帮我把老伴儿送到卫生院。」',
         '→ 信任不是从「签合同」开始，而是从「搭把手」开始。'),
        ("案例 02", "种粮大户李师傅",
         '「我认识你们行长，他去年帮我联系过收粮的渠道。」',
         "→ 一笔贷款的背后，可能是三顿饭、两次下田、一次牵线。"),
        ("案例 03", "独居老人张奶奶",
         '「你们这儿凉快，还有人说话，我每天都来坐会儿。」',
         '→ 老年客群的「第三空间」需求，比金融需求更早被唤醒。'),
    ]
    yy = Inches(2.1)
    for tag, person, quote, insight in cases:
        add_rect(s, Inches(0.8), yy, Inches(12), Inches(1.55), fill=WHITE, line=ORANGE_LIGHT)
        add_rect(s, Inches(0.8), yy, Inches(0.18), Inches(1.55), fill=ORANGE)
        add_text(s, Inches(1.2), yy + Inches(0.12), Inches(1.5), Inches(0.35),
                 tag, size=11, bold=True, color=ORANGE)
        add_text(s, Inches(1.2), yy + Inches(0.42), Inches(3), Inches(0.4),
                 person, size=15, bold=True, color=TEXT_DARK)
        add_text(s, Inches(4.3), yy + Inches(0.25), Inches(7.5), Inches(0.6),
                 quote, size=13, color=TEXT_DARK, font="楷体")
        add_rect(s, Inches(4.3), yy + Inches(0.92), Inches(7.5), Inches(0.5), fill=ORANGE_SOFT)
        add_text(s, Inches(4.5), yy + Inches(0.97), Inches(7.2), Inches(0.42),
                 insight, size=12, bold=True, color=ORANGE_DARK)
        yy += Inches(1.7)
    add_page_number(s, 7, TOTAL_SLIDES)


def slide_p2_relation():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SW, SH, fill=BG_CREAM)
    add_top_bar(s)
    add_bottom_bar(s)
    add_text(s, Inches(0.8), Inches(0.45), Inches(2), Inches(0.35),
             "PART 02 · 需求观察", size=12, bold=True, color=ORANGE)
    add_text(s, Inches(0.8), Inches(0.85), Inches(12), Inches(0.8),
             "流量思维 ≠ 关系思维：两种获客逻辑对比", size=28, bold=True, color=TEXT_DARK)
    add_rect(s, Inches(0.8), Inches(1.72), Inches(1), Inches(0.06), fill=ORANGE)
    headers = ["维度", "流量思维（城市网点）", "关系思维（乡镇网点）"]
    hx = [Inches(0.8), Inches(3.2), Inches(7.6)]
    hw = [Inches(2.3), Inches(4.3), Inches(4.9)]
    for i, h in enumerate(headers):
        fill = ORANGE if i > 0 else ORANGE_DARK
        add_rect(s, hx[i], Inches(2.1), hw[i], Inches(0.6), fill=fill)
        add_text(s, hx[i], Inches(2.17), hw[i], Inches(0.45),
                 h, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    rows = [
        ("获客来源", "线上广告 · 地推派单 · 厅堂流量", "老客转介绍 · 村干部引荐 · 熟人圈层"),
        ("信任建立", "靠品牌背书 + 合同条款", '靠「脸熟」 + 一次一次的真实帮忙'),
        ("决策周期", "短平快，当场比价", "长周期，需要三到五轮接触"),
        ("客群粘性", "低，利率敏感，容易搬家", "高，人情绑定，随时间增值"),
        ("成功指标", "进件数、批核率、单客 AUM", '转介绍率、村口的「知名度」'),
    ]
    ry = Inches(2.7)
    for r, (dim, left, right) in enumerate(rows):
        bg = WHITE if r % 2 == 0 else ORANGE_SOFT
        for i in range(3):
            add_rect(s, hx[i], ry, hw[i], Inches(0.78), fill=bg,
                     line=ORANGE_LIGHT)
        add_text(s, hx[0] + Inches(0.15), ry + Inches(0.2), hw[0] - Inches(0.3), Inches(0.4),
                 dim, size=12, bold=True, color=TEXT_DARK, align=PP_ALIGN.CENTER)
        add_text(s, hx[1] + Inches(0.2), ry + Inches(0.15), hw[1] - Inches(0.4), Inches(0.55),
                 left, size=11, color=TEXT_GRAY)
        add_text(s, hx[2] + Inches(0.2), ry + Inches(0.15), hw[2] - Inches(0.4), Inches(0.55),
                 right, size=11, bold=True, color=ORANGE_DARK)
        ry += Inches(0.78)
    add_rect(s, Inches(0.8), Inches(6.75), Inches(11.7), Inches(0.5), fill=ORANGE)
    add_text(s, Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.4),
             '💡  核心洞察：在乡镇，最大的流量不是路过的人，而是愿意「为你说话」的人。',
             size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_page_number(s, 8, TOTAL_SLIDES)


def slide_p3_direction():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SW, SH, fill=BG_CREAM)
    add_top_bar(s)
    add_bottom_bar(s)
    add_text(s, Inches(0.8), Inches(0.45), Inches(2), Inches(0.35),
             "PART 03 · 变革方向", size=12, bold=True, color=ORANGE)
    add_text(s, Inches(0.8), Inches(0.85), Inches(12), Inches(0.8),
             '什么是「小而实用」的社区中心？', size=28, bold=True, color=TEXT_DARK)
    add_rect(s, Inches(0.8), Inches(1.72), Inches(1), Inches(0.06), fill=ORANGE)
    cx, cy = Inches(6.65), Inches(4.6)
    R_big = Inches(2.0)
    add_rect(s, cx - R_big, cy - R_big, R_big * 2, R_big * 2, fill=ORANGE, shape=MSO_SHAPE.OVAL)
    add_text(s, cx - R_big, cy - Inches(0.9), R_big * 2, Inches(0.7),
             "社区中心", size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, cx - R_big, cy - Inches(0.2), R_big * 2, Inches(0.7),
             "Community Hub", size=14, color=ORANGE_LIGHT, align=PP_ALIGN.CENTER, font="Arial")
    nodes = [
        ("便民服务区", "代缴·打印·快递代收", Inches(1.8), Inches(2.2)),
        ("老年社交角", "歇脚·手机课堂·茶歇", Inches(11.5), Inches(2.2)),
        ("金融咨询台", "贷款·存款·理财咨询", Inches(1.8), Inches(6.8)),
        ("乡村信息站", "招工·农技·收粮信息", Inches(11.5), Inches(6.8)),
    ]
    import math
    for title, sub, nx, ny in nodes:
        dx = nx - cx
        dy = ny - cy
        dist = math.sqrt(dx**2 + dy**2)
        ux, uy = dx / dist, dy / dist
        lx1 = cx + ux * R_big
        ly1 = cy + uy * R_big
        lx2 = nx - ux * Inches(1.3)
        ly2 = ny - uy * Inches(0.55)
        connector = s.shapes.add_connector(1, lx1, ly1, lx2, ly2)
        connector.line.color.rgb = ORANGE
        connector.line.width = Pt(2)
        add_rect(s, nx - Inches(1.4), ny - Inches(0.7), Inches(2.8), Inches(1.4),
                 fill=WHITE, line=ORANGE)
        add_text(s, nx - Inches(1.3), ny - Inches(0.55), Inches(2.6), Inches(0.45),
                 title, size=14, bold=True, color=ORANGE_DARK, align=PP_ALIGN.CENTER)
        add_text(s, nx - Inches(1.3), ny - Inches(0.05), Inches(2.6), Inches(0.45),
                 sub, size=11, color=TEXT_GRAY, align=PP_ALIGN.CENTER)
    add_page_number(s, 10, TOTAL_SLIDES)


def slide_p3_actions():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SW, SH, fill=BG_CREAM)
    add_top_bar(s)
    add_bottom_bar(s)
    add_text(s, Inches(0.8), Inches(0.45), Inches(2), Inches(0.35),
             "PART 03 · 变革方向", size=12, bold=True, color=ORANGE)
    add_text(s, Inches(0.8), Inches(0.85), Inches(12), Inches(0.8),
             '落地建议：四件「低成本、高感知」的小事', size=28, bold=True, color=TEXT_DARK)
    add_rect(s, Inches(0.8), Inches(1.72), Inches(1), Inches(0.06), fill=ORANGE)
    actions = [
        ("①", '门口放一把「爱心伞架」',
         "下雨时村民随手借用，不用登记。伞面上印泰隆 logo = 流动的广告牌。",
         "成本：¥200 / 20 把", "插入：爱心伞 / 歇脚凳 照片"),
        ("②", '设一个「村口信息板」',
         "A3 纸打印招工、收粮、农技信息，每周一换。村民为了看板而进门。",
         "成本：¥10 / 周", "插入：信息板 / 公告栏 照片"),
        ("③", '每周三下午「手机课堂」',
         "教老人连 WiFi、视频通话、看医保。一场 10 个人 = 10 个家庭的关系锚点。",
         "成本：0（人力投入）", "插入：手机教学 / 培训场景"),
        ("④", '名片 + 一张「帮忙清单」',
         '除了联系方式，背面写「能帮你做的事」：代缴、代查、翻译、带路…',
         "成本：¥50 / 500 张", "插入：名片设计 / 物料示意"),
    ]
    positions = [
        (Inches(0.8), Inches(2.1)),
        (Inches(6.95), Inches(2.1)),
        (Inches(0.8), Inches(4.75)),
        (Inches(6.95), Inches(4.75)),
    ]
    for (num, title, desc, cost, img_lbl), (px, py) in zip(actions, positions):
        add_rect(s, px, py, Inches(5.85), Inches(2.55), fill=WHITE, line=ORANGE_LIGHT)
        add_rect(s, px, py, Inches(0.9), Inches(2.55), fill=ORANGE_LIGHT)
        add_text(s, px, py + Inches(0.8), Inches(0.9), Inches(1.0),
                 num, size=40, bold=True, color=ORANGE, align=PP_ALIGN.CENTER, font="Arial")
        add_text(s, px + Inches(1.05), py + Inches(0.15), Inches(4.6), Inches(0.45),
                 title, size=15, bold=True, color=TEXT_DARK)
        add_text(s, px + Inches(1.05), py + Inches(0.6), Inches(4.6), Inches(1.0),
                 desc, size=11, color=TEXT_GRAY)
        add_rect(s, px + Inches(1.05), py + Inches(1.6), Inches(4.6), Inches(0.3),
                 fill=ORANGE_SOFT)
        add_text(s, px + Inches(1.15), py + Inches(1.63), Inches(4.4), Inches(0.25),
                 "💰  %s" % cost, size=10, bold=True, color=ORANGE_DARK)
        lbl = img_lbl.split("：")[1] if "：" in img_lbl else img_lbl
        img_placeholder(s, px + Inches(3.9), py + Inches(1.95), Inches(1.85), Inches(0.5), lbl)
    add_page_number(s, 11, TOTAL_SLIDES)


def slide_p4_logic():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SW, SH, fill=BG_CREAM)
    add_top_bar(s)
    add_bottom_bar(s)
    add_text(s, Inches(0.8), Inches(0.45), Inches(2), Inches(0.35),
             "PART 04 · 核心逻辑", size=12, bold=True, color=ORANGE)
    add_text(s, Inches(0.8), Inches(0.85), Inches(12), Inches(0.8),
             '飞轮模型：非金融 → 人 → 关系 → 金融', size=28, bold=True, color=TEXT_DARK)
    add_rect(s, Inches(0.8), Inches(1.72), Inches(1), Inches(0.06), fill=ORANGE)
    steps = [
        ("LAYER 1", "非金融服务", "免费茶水 · 代缴 · 信息板 · 手机课",
         ORANGE_LIGHT, TEXT_DARK),
        ("LAYER 2", '带来「人」', "进门门槛 ≈ 0，村民愿意来、坐得住",
         ORANGE_SOFT, ORANGE_DARK),
        ("LAYER 3", '建立「关系」', '脸熟 + 帮忙 + 信任 = 从「客户经理」变「自己人」',
         ORANGE, WHITE),
        ("LAYER 4", '留下「金融」', "存款、贷款、理财，自然发生；转介绍滚雪球",
         ORANGE_DARK, WHITE),
    ]
    sy = Inches(2.1)
    for layer, title, sub, fill, color in steps:
        add_rect(s, Inches(0.8), sy, Inches(8), Inches(1.1), fill=fill)
        add_text(s, Inches(1.05), sy + Inches(0.1), Inches(2), Inches(0.4),
                 layer, size=11, bold=True, color=color, font="Arial")
        add_text(s, Inches(1.05), sy + Inches(0.42), Inches(7.5), Inches(0.55),
                 title, size=20, bold=True, color=color)
        add_text(s, Inches(8.7), sy + Inches(0.2), Inches(4), Inches(0.7),
                 sub, size=12, color=TEXT_GRAY, anchor=MSO_ANCHOR.MIDDLE)
        sy += Inches(1.18)
    arrow_x = Inches(8.4)
    for ay in [Inches(2.62), Inches(3.8), Inches(4.98)]:
        add_text(s, arrow_x, ay, Inches(0.3), Inches(0.5),
                 "⬇", size=18, color=ORANGE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Inches(0.8), Inches(6.95), Inches(11.7), Inches(0.38), fill=ORANGE_DARK)
    add_text(s, Inches(0.8), Inches(6.99), Inches(11.7), Inches(0.3),
             '一句话总结：先用非金融服务把「请进来」做轻，再用信任关系把「做业务」做重。',
             size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_page_number(s, 13, TOTAL_SLIDES)


def slide_p4_summary():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SW, SH, fill=BG_CREAM)
    add_top_bar(s)
    add_bottom_bar(s)
    add_text(s, Inches(0.8), Inches(0.45), Inches(3), Inches(0.35),
             "PART 04 · 核心逻辑", size=12, bold=True, color=ORANGE)
    add_text(s, Inches(0.8), Inches(0.85), Inches(12), Inches(0.8),
             "实习期间我的三点思考与成长", size=28, bold=True, color=TEXT_DARK)
    add_rect(s, Inches(0.8), Inches(1.72), Inches(1), Inches(0.06), fill=ORANGE)
    takeaways = [
        ("关于 业务",
         '小微金融的「风控」不只是看报表，更是看一个人在社区里的「口碑资产」。',
         "3 次", "陪同走访户数", "2 笔", "辅助促成贷款"),
        ("关于 沟通",
         '和村民说话要「去术语化」：不说「年化利率」，要说「借一万块，一天两块多利息」。',
         "7 篇", "撰写走访日志", "1 份", "整理话术手册"),
        ("关于 职业",
         '泰隆的「小微专营」不是一句口号，而是靠一个个网点、一位位经理站出来的。',
         "6 周", "全勤投入实习", "+∞", "对银行业的重新认知"),
    ]
    yy = Inches(2.1)
    for tag, insight, n1, l1, n2, l2 in takeaways:
        add_rect(s, Inches(0.8), yy, Inches(12), Inches(1.55), fill=WHITE, line=ORANGE_LIGHT)
        tag_parts = tag.split()
        add_rect(s, Inches(0.8), yy, Inches(1.8), Inches(1.55), fill=ORANGE)
        add_text(s, Inches(0.8), yy + Inches(0.35), Inches(1.8), Inches(0.4),
                 tag_parts[0], size=12, bold=True, color=ORANGE_LIGHT, align=PP_ALIGN.CENTER, font="Arial")
        add_text(s, Inches(0.8), yy + Inches(0.7), Inches(1.8), Inches(0.6),
                 tag_parts[1], size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text(s, Inches(2.85), yy + Inches(0.2), Inches(5.3), Inches(1.1),
                 insight, size=13, color=TEXT_DARK, anchor=MSO_ANCHOR.MIDDLE)
        add_rect(s, Inches(8.3), yy + Inches(0.2), Inches(2.1), Inches(1.15), fill=ORANGE_SOFT)
        add_rect(s, Inches(10.6), yy + Inches(0.2), Inches(2.1), Inches(1.15), fill=ORANGE_LIGHT)
        add_text(s, Inches(8.3), yy + Inches(0.3), Inches(2.1), Inches(0.55),
                 n1, size=22, bold=True, color=ORANGE_DARK, align=PP_ALIGN.CENTER, font="Arial")
        add_text(s, Inches(8.3), yy + Inches(0.85), Inches(2.1), Inches(0.4),
                 l1, size=10, color=TEXT_GRAY, align=PP_ALIGN.CENTER)
        add_text(s, Inches(10.6), yy + Inches(0.3), Inches(2.1), Inches(0.55),
                 n2, size=22, bold=True, color=ORANGE_DARK, align=PP_ALIGN.CENTER, font="Arial")
        add_text(s, Inches(10.6), yy + Inches(0.85), Inches(2.1), Inches(0.4),
                 l2, size=10, color=TEXT_GRAY, align=PP_ALIGN.CENTER)
        yy += Inches(1.7)
    add_page_number(s, 14, TOTAL_SLIDES)


def slide_end():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SW, SH, fill=ORANGE)
    add_rect(s, 0, Inches(6.9), SW, Inches(0.6), fill=ORANGE_DARK)
    add_rect(s, Inches(-0.5), Inches(5.5), Inches(3.5), Inches(3.5), fill=ORANGE_LIGHT, shape=MSO_SHAPE.OVAL)
    add_rect(s, Inches(10.8), Inches(-0.8), Inches(3.8), Inches(3.8), fill=ORANGE_DARK, shape=MSO_SHAPE.OVAL)
    add_rect(s, Inches(11.5), Inches(5.0), Inches(2.2), Inches(2.2), fill=WHITE, shape=MSO_SHAPE.OVAL)
    add_rect(s, Inches(10.2), Inches(6.0), Inches(1.0), Inches(1.0), fill=ORANGE_LIGHT, shape=MSO_SHAPE.OVAL)
    add_text(s, Inches(0.8), Inches(1.2), Inches(12), Inches(0.5),
             "THANK YOU", size=72, bold=True, color=WHITE, align=PP_ALIGN.CENTER, font="Arial")
    add_text(s, Inches(0.8), Inches(2.8), Inches(12), Inches(0.8),
             "感谢聆听 · 敬请指正", size=32, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_rect(s, Inches(6.1), Inches(3.85), Inches(1.1), Inches(0.06), fill=WHITE)
    add_text(s, Inches(0.8), Inches(4.2), Inches(12), Inches(0.45),
             "答辩人：曾纪元   |   泰隆银行小微专营实习项目",
             size=15, color=ORANGE_LIGHT, align=PP_ALIGN.CENTER)
    contact = [
        ("📧", "zeng.ian@outlook.com"),
        ("📱", "+86 136 8791 8205"),
        ("🏫", "江西财经大学 · 27 届"),
    ]
    cx0 = Inches(3.4)
    for i, (ic, val) in enumerate(contact):
        add_rect(s, cx0 + Inches(i * 2.2), Inches(5.0), Inches(2.0), Inches(0.7),
                 fill=ORANGE_DARK, line=ORANGE_LIGHT)
        add_text(s, cx0 + Inches(i * 2.2) + Inches(0.1), Inches(5.1), Inches(0.5), Inches(0.5),
                 ic, size=18, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, cx0 + Inches(i * 2.2) + Inches(0.55), Inches(5.1), Inches(1.35), Inches(0.5),
                 val, size=12, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(0.8), Inches(7.0), Inches(12), Inches(0.35),
             "TAILONG BANK · INTERNSHIP DEFENSE 2026",
             size=10, color=ORANGE_LIGHT, align=PP_ALIGN.CENTER, font="Arial")


def build():
    slide_cover()
    slide_toc()
    slide_part_intro(1, "网点现状", "网点的角色，正在悄悄发生变化。",
                     '银行的围墙，正在被一杯免费的茶水慢慢拆掉。', 3)
    slide_p1_status()
    slide_p1_shift()
    slide_part_intro(2, "需求观察", '在乡镇做金融，首先要读懂什么是「关系」。',
                     '一笔贷款的背后，可能是三顿饭、两次下田、一次牵线。', 6)
    slide_p2_observation()
    slide_p2_relation()
    slide_part_intro(3, "变革方向", '把网点做「小」，但把服务做「实」。',
                     '不用花大钱装修，只要四件「低成本、高感知」的小事。', 9)
    slide_p3_direction()
    slide_p3_actions()
    slide_part_intro(4, "核心逻辑", '非金融服务是「引子」，金融服务是「结果」。',
                     '请进来 → 坐得住 → 信得过 → 成业务。这就是小微专营的乡镇公式。', 12)
    slide_p4_logic()
    slide_p4_summary()
    slide_end()
    import shutil
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), OUTPUT_LOCAL)
    prs.save(path)
    print("OK PPT saved locally: %s" % path)
    print("Total slides: %d" % len(prs.slides))
    try:
        shutil.copy2(path, OUTPUT_TARGET)
        print("Copied to target: %s" % OUTPUT_TARGET)
    except Exception as e:
        print("(Skip copy to target: %s)" % e)
        print("Please copy manually from: %s" % path)


if __name__ == "__main__":
    build()
